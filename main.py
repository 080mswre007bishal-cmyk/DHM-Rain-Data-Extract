"""
main.py
-------

For every station in stations.json:

  1. Go to the DHM rainfall-watch-map page.
  2. In the "SORT & FILTERS" panel, set Filter type to
     "Station No." and select this station's Station Index.
  3. Click "Apply Filter".
  4. Click the station marker.
  5. In the station popup/panel, select "Daily".
  6. Click "VIEW".
  7. Read the Date | Daily Rainfall table.
  8. Skip today's row.
  9. Save the data locally as output/<station>.xlsx.
 10. Create/find a separate Google Spreadsheet for this station.
 11. Add only dates that are not already present in that
     station's Google Spreadsheet.

Google Drive structure:

    DHM Rainfall/
        1001_Dipayal
        1002_Tarakhola
        1003_Panchase
        ...

Each station has its own Google Spreadsheet.

Each station spreadsheet contains:

    Date | Daily Rainfall (mm)

There are NO station folders.

Duplicate checking is performed by DATE because each
Google Spreadsheet belongs to only one station.

Browser mode:
    Local Windows  -> headed
    GitHub Actions -> headless
"""

import datetime
import json
import os
import re
import time

from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright

from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
from googleapiclient.discovery import build


# ============================================================
# CONFIGURATION
# ============================================================

STATIONS_FILE = "stations.json"

OUTPUT_DIR = "output"
DEBUG_DIR = os.path.join(
    OUTPUT_DIR,
    "debug"
)

MAP_URL = (
    "https://dhm.gov.np/"
    "hydrology/rainfall-watch-map"
)

# Existing Google Drive folder.
# All station spreadsheets go DIRECTLY inside this folder.
GOOGLE_FOLDER_NAME = "DHM Rainfall"

# Worksheet/tab inside every station spreadsheet.
GOOGLE_SHEET_TAB = "Rainfall"

# Google OAuth scopes.
GOOGLE_SCOPES = [
    "https://www.googleapis.com/auth/drive",
    "https://www.googleapis.com/auth/spreadsheets",
]


# ============================================================
# DEBUG STATION CONTROL
# ============================================================

# DEBUG_STATIONS=1
#     -> process only first station
#
# DEBUG_STATIONS=5
#     -> process first five stations
#
# DEBUG_STATIONS=0
#     -> process all stations

DEBUG_STATIONS = int(
    os.environ.get(
        "DEBUG_STATIONS",
        "0"
    )
)


# ============================================================
# NEPAL TIME
# ============================================================

NEPAL_OFFSET = datetime.timedelta(
    hours=5,
    minutes=45
)


# ============================================================
# DATE PARSING
# ============================================================

DATE_FORMATS = [
    "%a, %b %d, %Y, %I:%M %p",
    "%a, %b %d, %Y",
    "%Y-%m-%d",
    "%d %b, %Y",
    "%d %B, %Y",
    "%d %b %Y",
    "%d %B %Y",
    "%d/%m/%Y",
    "%m/%d/%Y",
    "%d-%m-%Y",
]


DATE_PATTERNS = [
    r"[A-Za-z]{3},\s+[A-Za-z]{3}\s+\d{1,2},\s+\d{4}(?:,\s+\d{1,2}:\d{2}\s*[AP]M)?",
    r"\d{4}-\d{2}-\d{2}",
    r"\d{1,2}\s+[A-Za-z]{3,9},?\s+\d{4}",
    r"\d{1,2}/\d{1,2}/\d{4}",
    r"\d{1,2}-\d{1,2}-\d{4}",
]


# ============================================================
# FILE / NAME HELPERS
# ============================================================

def safe_filename(
    name,
    station_index,
    url_id
):

    base = (
        station_index
        or url_id
    )

    clean_name = re.sub(
        r"[^\w\-]+",
        "_",
        name
    ).strip("_")

    return (
        f"{base}_{clean_name}.xlsx"
    )


def safe_google_name(
    name,
    station_index,
    url_id
):
    """
    Creates the individual Google Spreadsheet name.

    Example:
        1009_Chautara
    """

    base = (
        str(station_index).strip()
        if station_index
        else str(url_id).strip()
    )

    clean_name = re.sub(
        r"[^\w\-]+",
        "_",
        str(name)
    ).strip("_")

    if clean_name:

        return (
            f"{base}_{clean_name}"
        )

    return base


def nepal_today_iso():

    now_utc = datetime.datetime.utcnow()

    now_npt = (
        now_utc
        + NEPAL_OFFSET
    )

    return now_npt.date().strftime(
        "%Y-%m-%d"
    )


# ============================================================
# DATE / NUMBER PARSING
# ============================================================

def parse_date_token(token):

    token = token.strip()

    for fmt in DATE_FORMATS:

        try:

            return datetime.datetime.strptime(
                token,
                fmt
            ).strftime(
                "%Y-%m-%d"
            )

        except ValueError:

            continue

    return None


def parse_number(text):

    if text is None:

        return None

    text = str(text).strip()

    # Remove commas from numbers.
    text = text.replace(
        ",",
        ""
    )

    # Handle common missing-value indicators.
    if text.lower() in {
        "",
        "-",
        "--",
        "n/a",
        "na",
        "null",
        "none",
    }:

        return None

    matches = re.findall(
        r"\d+\.\d+|\d+",
        text
    )

    if not matches:

        return None

    try:

        val = float(
            matches[0]
        )

    except ValueError:

        return None

    return (
        val
        if 0 <= val < 1000
        else None
    )


# ============================================================
# DEBUG HELPERS
# ============================================================

def save_debug_file(
    tag,
    suffix,
    content
):

    os.makedirs(
        DEBUG_DIR,
        exist_ok=True
    )

    safe = re.sub(
        r"[^\w.-]+",
        "_",
        tag
    )

    with open(
        os.path.join(
            DEBUG_DIR,
            f"{safe}{suffix}"
        ),
        "w",
        encoding="utf-8",
    ) as f:

        f.write(
            str(content)
        )


# ============================================================
# STATION FILTER
# ============================================================

def open_station_via_filter(
    page,
    station_index,
    debug_tag=None
):
    """
    DHM uses a two-stage filter:

    1. Select "Station No.".
    2. Wait for the Station No. dropdown.
    3. Select the requested station number.
    4. Click Apply Filter.
    """

    filter_type_ok = False
    station_number_ok = False
    apply_ok = False

    station_index = str(
        station_index
    ).strip()

    # --------------------------------------------------------
    # STEP 1: Select Station No.
    # --------------------------------------------------------

    try:

        selects = page.locator(
            "select"
        )

        for i in range(
            selects.count()
        ):

            sel = selects.nth(i)

            try:

                opts = [
                    o.strip()
                    for o in sel.locator(
                        "option"
                    ).all_inner_texts()
                ]

            except Exception:

                continue

            matching_option = None

            for option in opts:

                if (
                    "station no"
                    in option.lower()
                ):

                    matching_option = option

                    break

            if matching_option:

                print(
                    f"   selecting Filter Type: "
                    f"{matching_option}"
                )

                sel.select_option(
                    label=matching_option,
                    timeout=3000,
                )

                filter_type_ok = True

                break

    except Exception as e:

        print(
            f"   Filter Type selection error: "
            f"{e}"
        )

    if not filter_type_ok:

        print(
            "   Could not find/select "
            "'Station No.' in Filter Type."
        )

        return False

    # --------------------------------------------------------
    # STEP 2: Wait for station number dropdown
    # --------------------------------------------------------

    print(
        f"   waiting for Station No. dropdown "
        f"containing {station_index}..."
    )

    for attempt in range(30):

        try:

            page.wait_for_timeout(
                500
            )

            selects = page.locator(
                "select"
            )

            for i in range(
                selects.count()
            ):

                sel = selects.nth(i)

                try:

                    opts = [
                        o.strip()
                        for o in sel.locator(
                            "option"
                        ).all_inner_texts()
                    ]

                except Exception:

                    continue

                if station_index in opts:

                    print(
                        f"   found Station No. "
                        f"dropdown containing "
                        f"{station_index}"
                    )

                    sel.select_option(
                        label=station_index,
                        timeout=3000,
                    )

                    station_number_ok = True

                    print(
                        f"   selected Station No.: "
                        f"{station_index}"
                    )

                    break

            if station_number_ok:

                break

        except Exception as e:

            if attempt == 29:

                print(
                    f"   Station No. selection "
                    f"error: {e}"
                )

    if not station_number_ok:

        print(
            f"   Station number "
            f"{station_index} was not found "
            f"in the dropdown."
        )

        if debug_tag:

            try:

                save_debug_file(
                    debug_tag,
                    "_station_filter_debug.txt",
                    page.locator(
                        "body"
                    ).inner_text(),
                )

            except Exception:

                pass

        return False

    # --------------------------------------------------------
    # STEP 3: Click Apply Filter
    # --------------------------------------------------------

    print(
        "   clicking Apply Filter..."
    )

    for selector in [
        "button:has-text('Apply Filter')",
        "text='Apply Filter'",
        "input[value='Apply Filter']",
    ]:

        try:

            button = page.locator(
                selector
            ).first

            if button.is_visible():

                button.click(
                    timeout=5000
                )

                apply_ok = True

                print(
                    "   Apply Filter clicked."
                )

                break

        except Exception:

            continue

    if not apply_ok:

        print(
            "   Could not click Apply Filter."
        )

    # --------------------------------------------------------
    # Wait for filtered map
    # --------------------------------------------------------

    page.wait_for_timeout(
        2000
    )

    # --------------------------------------------------------
    # Debug screenshot
    # --------------------------------------------------------

    if debug_tag:

        try:

            safe_debug_tag = re.sub(
                r"[^\w.-]+",
                "_",
                debug_tag,
            )

            page.screenshot(
                path=os.path.join(
                    DEBUG_DIR,
                    f"{safe_debug_tag}_after_filter.png",
                )
            )

        except Exception:

            pass

    return (
        filter_type_ok
        and station_number_ok
        and apply_ok
    )


# ============================================================
# MAP MARKER
# ============================================================

def get_marker_locator(page):

    for selector in [
        "#map .leaflet-marker-icon",
        ".leaflet-container .leaflet-marker-icon",
    ]:

        try:

            loc = page.locator(
                selector
            )

            if loc.count() > 0:

                return loc

        except Exception:

            pass

    return None


# ============================================================
# STATION POPUP / PANEL
# ============================================================

def _candidate_panels(page):

    selectors = [
        "[role='dialog']:visible",
        ".modal.show:visible",
        ".modal.in:visible",
        ".offcanvas.show:visible",
        "div[class*='station']:visible",
        "section[class*='station']:visible",
        "div[class*='popup']:visible",
        "section[class*='popup']:visible",
    ]

    result = []

    for selector in selectors:

        try:

            loc = page.locator(
                selector
            )

            for i in range(
                loc.count()
            ):

                el = loc.nth(i)

                try:

                    if not el.is_visible():

                        continue

                    inside_map = el.evaluate(
                        "el => !!el.closest('.leaflet-container')"
                    )

                    if inside_map:

                        continue

                    txt = el.inner_text(
                        timeout=1000
                    )

                    lower_txt = txt.lower()

                    if (
                        "station index"
                        in lower_txt
                        and "view"
                        in lower_txt
                    ):

                        result.append(
                            el
                        )

                except Exception:

                    pass

        except Exception:

            pass

    return result


def find_station_panel(page):

    panels = _candidate_panels(
        page
    )

    if not panels:

        return None

    # Prefer the smallest matching visible panel.
    ranked = sorted(
        panels,
        key=lambda p: len(
            p.inner_text()
        ),
    )

    return ranked[0]


# ============================================================
# DAILY PERIOD
# ============================================================

def set_period_to_daily_in_panel(
    panel
):
    """
    Select Daily in the Period dropdown.
    """

    print(
        "   selecting Period: Daily..."
    )

    selects = panel.locator(
        "select"
    )

    for i in range(
        selects.count()
    ):

        sel = selects.nth(i)

        try:

            opts = [
                o.strip()
                for o in sel.locator(
                    "option"
                ).all_inner_texts()
            ]

            daily_option = None

            for option in opts:

                if option.lower() == "daily":

                    daily_option = option

                    break

            if daily_option:

                sel.select_option(
                    label=daily_option,
                    timeout=3000,
                )

                print(
                    "   Daily selected."
                )

                return True

        except Exception:

            continue

    print(
        "   Could not find Daily option."
    )

    return False


# ============================================================
# VIEW BUTTON
# ============================================================

def click_view_in_panel(
    panel
):

    print(
        "   clicking VIEW..."
    )

    for selector in [
        "button:has-text('VIEW')",
        "input[value='VIEW']",
        "text=VIEW",
    ]:

        try:

            button = panel.locator(
                selector
            ).first

            if not button.is_visible():

                continue

            button.click(
                timeout=5000
            )

            print(
                "   VIEW clicked."
            )

            return True

        except Exception:

            continue

    print(
        "   Could not find/click VIEW."
    )

    return False


# ============================================================
# TABLE PARSING
# ============================================================

def table_rows_from_panel(
    panel
):
    """
    Read the DHM rainfall table.

    Returns:

        [
            (date_str, daily_rainfall),
            ...
        ]

    Running Total is intentionally ignored.

    The parser searches each row for a date instead of assuming
    that the date is always the first cell.
    """

    rows_data = []

    try:

        tables = panel.locator(
            "table"
        )

        table_count = tables.count()

        print(
            f"   Found {table_count} table(s) "
            f"inside station panel."
        )

    except Exception as e:

        print(
            f"   Could not locate tables: {e}"
        )

        return rows_data

    # --------------------------------------------------------
    # Examine every table
    # --------------------------------------------------------

    for table_index in range(
        table_count
    ):

        try:

            table = tables.nth(
                table_index
            )

            rows = table.locator(
                "tr"
            )

            row_count = rows.count()

            print(
                f"   Table {table_index + 1}: "
                f"{row_count} row(s)"
            )

        except Exception:

            continue

        # ----------------------------------------------------
        # Examine every row
        # ----------------------------------------------------

        for row_index in range(
            row_count
        ):

            try:

                row = rows.nth(
                    row_index
                )

                cells = row.locator(
                    "td, th"
                )

                cell_count = cells.count()

                texts = []

                for cell_index in range(
                    cell_count
                ):

                    try:

                        text = (
                            cells.nth(
                                cell_index
                            )
                            .inner_text()
                            .strip()
                        )

                    except Exception:

                        text = ""

                    texts.append(
                        text
                    )

            except Exception:

                continue

            if not texts:

                continue

            # ------------------------------------------------
            # Print first few rows for debugging
            # ------------------------------------------------

            if row_index < 5:

                print(
                    f"      row {row_index}: "
                    f"{texts}"
                )

            # ------------------------------------------------
            # Skip obvious header rows
            # ------------------------------------------------

            joined = " | ".join(
                texts
            )

            lower_joined = joined.lower()

            if (
                "date" in lower_joined
                and (
                    "daily" in lower_joined
                    or "rainfall" in lower_joined
                )
            ):

                continue

            # ------------------------------------------------
            # Find date anywhere in row
            # ------------------------------------------------

            date_str = None
            date_cell_index = None

            for cell_index, cell_text in enumerate(
                texts
            ):

                if not cell_text:

                    continue

                for pattern in DATE_PATTERNS:

                    match = re.search(
                        pattern,
                        cell_text,
                    )

                    if match:

                        parsed = parse_date_token(
                            match.group(0)
                        )

                        if parsed:

                            date_str = parsed
                            date_cell_index = (
                                cell_index
                            )

                            break

                if date_str:

                    break

            if not date_str:

                continue

            # ------------------------------------------------
            # Find rainfall value after date
            # ------------------------------------------------

            daily_val = None

            search_start = (
                date_cell_index + 1
                if date_cell_index is not None
                else 1
            )

            for cell_text in texts[
                search_start:
            ]:

                if not cell_text:

                    continue

                # Do not interpret another date as rainfall.
                contains_date = False

                for pattern in DATE_PATTERNS:

                    if re.search(
                        pattern,
                        cell_text
                    ):

                        contains_date = True

                        break

                if contains_date:

                    continue

                daily_val = parse_number(
                    cell_text
                )

                if daily_val is not None:

                    break

            if daily_val is None:

                continue

            rows_data.append(
                (
                    date_str,
                    daily_val,
                )
            )

    # --------------------------------------------------------
    # Remove duplicate dates
    # --------------------------------------------------------

    unique_rows = {}

    for (
        date_str,
        daily_val
    ) in rows_data:

        unique_rows[
            date_str
        ] = daily_val

    rows_data = sorted(
        unique_rows.items()
    )

    return rows_data


# ============================================================
# CLOSE PANEL
# ============================================================

def close_any_panel(page):

    for selector in [
        "[aria-label='Close']:visible",
        "button[title='Close']:visible",
        ".btn-close:visible",
        ".close:visible",
        "button:has-text('Close'):visible",
    ]:

        try:

            loc = page.locator(
                selector
            ).last

            if loc.is_visible():

                loc.click(
                    force=True,
                    timeout=1000,
                )

                page.wait_for_timeout(
                    250
                )

                return

        except Exception:

            pass

    try:

        page.keyboard.press(
            "Escape"
        )

        page.wait_for_timeout(
            250
        )

    except Exception:

        pass


# ============================================================
# PROCESS ONE STATION
# ============================================================

def process_station(
    page,
    station,
    debug_tag=None
):
    """
    Returns:

        [
            (date_str, daily_rainfall),
            ...
        ]

    Today's row is excluded.
    """

    station_index = station.get(
        "station_index"
    )

    if not station_index:

        raise RuntimeError(
            "This station has no station_index "
            "in stations.json."
        )

    # --------------------------------------------------------
    # Open map
    # --------------------------------------------------------

    print(
        "   Opening DHM rainfall map..."
    )

    page.goto(
        MAP_URL,
        wait_until="domcontentloaded",
        timeout=30000,
    )

    page.wait_for_timeout(
        2500
    )

    # --------------------------------------------------------
    # Apply station filter
    # --------------------------------------------------------

    filtered_ok = (
        open_station_via_filter(
            page,
            station_index,
            debug_tag=debug_tag,
        )
    )

    if not filtered_ok:

        raise RuntimeError(
            f"Could not apply the Station No. "
            f"filter for index "
            f"{station_index!r}."
        )

    # --------------------------------------------------------
    # Find station marker and click it
    # --------------------------------------------------------

    print(
        "   Looking for station marker..."
    )

    panel = None
    marker_clicked = False

    for attempt in range(
        30
    ):

        # First see whether panel already exists.
        panel = find_station_panel(
            page
        )

        if panel is not None:

            print(
                "   Station panel found."
            )

            break

        markers = get_marker_locator(
            page
        )

        if (
            markers
            and markers.count() > 0
        ):

            try:

                marker_count = markers.count()

                print(
                    f"   Found {marker_count} "
                    f"map marker(s)."
                )

                # After filtering, the first visible marker
                # should correspond to the selected station.
                markers.first.click(
                    force=True,
                    timeout=3000,
                )

                marker_clicked = True

                print(
                    "   Station marker clicked."
                )

                page.wait_for_timeout(
                    1000
                )

            except Exception as e:

                if attempt in {
                    0,
                    5,
                    10,
                    20,
                }:

                    print(
                        f"   Marker click attempt "
                        f"{attempt + 1} failed: "
                        f"{e}"
                    )

        time.sleep(
            0.5
        )

    # --------------------------------------------------------
    # No panel
    # --------------------------------------------------------

    if panel is None:

        if debug_tag:

            try:

                safe_debug_tag = re.sub(
                    r"[^\w.-]+",
                    "_",
                    debug_tag
                )

                page.screenshot(
                    path=os.path.join(
                        DEBUG_DIR,
                        f"{safe_debug_tag}_no_panel.png"
                    )
                )

                save_debug_file(
                    debug_tag,
                    "_no_panel_text.txt",
                    page.locator(
                        "body"
                    ).inner_text(),
                )

                save_debug_file(
                    debug_tag,
                    "_no_panel_html.txt",
                    page.content(),
                )

            except Exception:

                pass

        raise RuntimeError(
            "Filtered successfully but no station "
            "popup panel appeared."
        )

    # --------------------------------------------------------
    # Save panel debug files
    # --------------------------------------------------------

    if debug_tag:

        try:

            safe_debug_tag = re.sub(
                r"[^\w.-]+",
                "_",
                debug_tag,
            )

            panel.screenshot(
                path=os.path.join(
                    DEBUG_DIR,
                    f"{safe_debug_tag}_panel.png"
                )
            )

        except Exception:

            pass

        try:

            save_debug_file(
                debug_tag,
                "_panel_html.txt",
                panel.evaluate(
                    "el => el.outerHTML"
                ),
            )

        except Exception:

            pass

        try:

            save_debug_file(
                debug_tag,
                "_panel_text.txt",
                panel.inner_text(),
            )

        except Exception:

            pass

    # --------------------------------------------------------
    # Select Daily
    # --------------------------------------------------------

    period_ok = (
        set_period_to_daily_in_panel(
            panel
        )
    )

    if not period_ok:

        print(
            "   WARNING: Daily period could "
            "not be selected."
        )

    # --------------------------------------------------------
    # Click VIEW
    # --------------------------------------------------------

    view_ok = (
        click_view_in_panel(
            panel
        )
    )

    if not view_ok:

        print(
            "   WARNING: VIEW button could "
            "not be clicked."
        )

    # --------------------------------------------------------
    # Wait for rainfall table
    # --------------------------------------------------------

    print(
        "   Waiting for rainfall table..."
    )

    try:

        panel.locator(
            "table"
        ).first.wait_for(
            state="visible",
            timeout=10000,
        )

        print(
            "   Rainfall table became visible."
        )

    except Exception:

        print(
            "   Warning: rainfall table did not "
            "become visible within 10 seconds."
        )

    # Give the DHM JavaScript time to populate rows.
    page.wait_for_timeout(
        3000
    )

    # --------------------------------------------------------
    # Read table
    # --------------------------------------------------------

    all_rows = (
        table_rows_from_panel(
            panel
        )
    )

    # --------------------------------------------------------
    # Remove today's row
    # --------------------------------------------------------

    today = nepal_today_iso()

    kept_rows = [
        row
        for row in all_rows
        if row[0] != today
    ]

    if today in {
        row[0]
        for row in all_rows
    }:

        print(
            f"   Skipped today's row: {today}"
        )

    # --------------------------------------------------------
    # Debug result
    # --------------------------------------------------------

    if debug_tag:

        save_debug_file(
            debug_tag,
            "_result.txt",
            f"filtered_ok: {filtered_ok}\n"
            f"marker_clicked: {marker_clicked}\n"
            f"period_ok: {period_ok}\n"
            f"view_ok: {view_ok}\n"
            f"today: {today}\n"
            f"all_rows: {all_rows}\n"
            f"kept_rows: {kept_rows}\n",
        )

    # --------------------------------------------------------
    # Close popup
    # --------------------------------------------------------

    close_any_panel(
        page
    )

    return kept_rows


# ============================================================
# LOCAL EXCEL OUTPUT
# ============================================================

def append_all_rows(
    path,
    entries
):
    """
    Local Excel output.

    Columns:

        Date
        Daily Rainfall (mm)

    Existing dates are checked so local Excel files also
    do not accumulate duplicate rows.
    """

    existing_dates = set()

    if os.path.exists(path):

        wb = load_workbook(
            path
        )

        ws = wb.active

        # ----------------------------------------------------
        # Read existing dates
        # ----------------------------------------------------

        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            if not row:

                continue

            value = row[0]

            if value is None:

                continue

            if isinstance(
                value,
                datetime.datetime
            ):

                date_str = (
                    value.date()
                    .strftime(
                        "%Y-%m-%d"
                    )
                )

            elif isinstance(
                value,
                datetime.date
            ):

                date_str = (
                    value.strftime(
                        "%Y-%m-%d"
                    )
                )

            else:

                date_str = str(
                    value
                ).strip()

                parsed = parse_date_token(
                    date_str
                )

                if parsed:

                    date_str = parsed

            existing_dates.add(
                date_str
            )

    else:

        wb = Workbook()

        ws = wb.active

        ws.title = "Rainfall"

        ws.append(
            [
                "Date",
                "Daily Rainfall (mm)",
            ]
        )

    # --------------------------------------------------------
    # Append only new dates
    # --------------------------------------------------------

    added = 0

    for (
        date_str,
        daily_val
    ) in sorted(entries):

        if date_str in existing_dates:

            continue

        ws.append(
            [
                date_str,
                daily_val,
            ]
        )

        existing_dates.add(
            date_str
        )

        added += 1

    wb.save(
        path
    )

    return added


# ============================================================
# GOOGLE AUTHENTICATION
# ============================================================

def load_google_credentials():
    """
    Load OAuth credentials from token.json.

    GitHub Actions creates token.json from the
    GOOGLE_TOKEN_JSON repository secret.
    """

    token_path = "token.json"

    if not os.path.exists(
        token_path
    ):

        raise FileNotFoundError(
            "token.json was not found. "
            "Make sure the GitHub workflow restores "
            "GOOGLE_TOKEN_JSON before running main.py."
        )

    creds = (
        Credentials.from_authorized_user_file(
            token_path,
            GOOGLE_SCOPES
        )
    )

    if creds.expired:

        if creds.refresh_token:

            print(
                "Google access token expired. "
                "Refreshing..."
            )

            creds.refresh(
                Request()
            )

        else:

            raise RuntimeError(
                "Google credentials have expired "
                "and no refresh token is available."
            )

    print(
        "Google credentials loaded."
    )

    return creds


# ============================================================
# GOOGLE DRIVE FOLDER
# ============================================================

def get_google_drive_folder(
    drive,
    folder_name
):
    """
    Find the existing DHM Rainfall folder.

    The folder must already exist in Google Drive.
    """

    print(
        f"Searching for Google Drive folder: "
        f"{folder_name}"
    )

    escaped_name = (
        folder_name.replace(
            "'",
            "\\'"
        )
    )

    query = (
        f"name = '{escaped_name}' "
        f"and mimeType = "
        f"'application/vnd.google-apps.folder' "
        f"and trashed = false"
    )

    results = drive.files().list(
        q=query,
        pageSize=100,
        fields="files(id,name)"
    ).execute()

    files = results.get(
        "files",
        []
    )

    if not files:

        raise RuntimeError(
            f"Could not find Google Drive folder: "
            f"{folder_name}"
        )

    folder = files[0]

    print(
        f"Google Drive folder found: "
        f"{folder['name']}"
    )

    print(
        f"Folder ID: {folder['id']}"
    )

    return folder["id"]


# ============================================================
# GOOGLE SPREADSHEET FIND
# ============================================================

def find_station_spreadsheet(
    drive,
    folder_id,
    spreadsheet_name
):
    """
    Find a station spreadsheet specifically inside
    the DHM Rainfall folder.

    There are no station subfolders.
    """

    escaped_name = (
        spreadsheet_name.replace(
            "'",
            "\\'"
        )
    )

    query = (
        f"name = '{escaped_name}' "
        f"and mimeType = "
        f"'application/vnd.google-apps.spreadsheet' "
        f"and '{folder_id}' in parents "
        f"and trashed = false"
    )

    results = drive.files().list(
        q=query,
        pageSize=10,
        fields="files(id,name,parents)"
    ).execute()

    files = results.get(
        "files",
        []
    )

    if files:

        return files[0]["id"]

    return None


# ============================================================
# GOOGLE SPREADSHEET CREATE
# ============================================================

def create_station_spreadsheet(
    drive,
    sheets,
    folder_id,
    spreadsheet_name
):
    """
    Create a new Google Spreadsheet for one station
    and place it directly inside DHM Rainfall.
    """

    print(
        f"   Creating Google Spreadsheet: "
        f"{spreadsheet_name}"
    )

    body = {
        "properties": {
            "title": spreadsheet_name
        },
        "sheets": [
            {
                "properties": {
                    "title": GOOGLE_SHEET_TAB
                }
            }
        ]
    }

    created = (
        sheets.spreadsheets()
        .create(
            body=body
        )
        .execute()
    )

    spreadsheet_id = created[
        "spreadsheetId"
    ]

    # --------------------------------------------------------
    # Put spreadsheet directly inside DHM Rainfall.
    # --------------------------------------------------------

    drive.files().update(
        fileId=spreadsheet_id,
        addParents=folder_id,
        fields="id,parents"
    ).execute()

    # --------------------------------------------------------
    # Add header
    # --------------------------------------------------------

    sheets.spreadsheets().values().update(
        spreadsheetId=spreadsheet_id,
        range=f"{GOOGLE_SHEET_TAB}!A1:B1",
        valueInputOption="USER_ENTERED",
        body={
            "values": [
                [
                    "Date",
                    "Daily Rainfall (mm)"
                ]
            ]
        }
    ).execute()

    print(
        "   Google Spreadsheet created."
    )

    print(
        f"   Spreadsheet ID: "
        f"{spreadsheet_id}"
    )

    return spreadsheet_id


# ============================================================
# GOOGLE SPREADSHEET GET OR CREATE
# ============================================================

def get_or_create_station_spreadsheet(
    drive,
    sheets,
    folder_id,
    station
):

    station_index = str(
        station.get(
            "station_index",
            ""
        )
    ).strip()

    name = str(
        station.get(
            "name",
            ""
        )
    ).strip()

    url_id = str(
        station.get(
            "url_id",
            ""
        )
    ).strip()

    spreadsheet_name = safe_google_name(
        name,
        station_index,
        url_id
    )

    spreadsheet_id = (
        find_station_spreadsheet(
            drive,
            folder_id,
            spreadsheet_name
        )
    )

    if spreadsheet_id:

        print(
            f"   Google Spreadsheet found: "
            f"{spreadsheet_name}"
        )

        return (
            spreadsheet_id,
            spreadsheet_name
        )

    spreadsheet_id = (
        create_station_spreadsheet(
            drive,
            sheets,
            folder_id,
            spreadsheet_name
        )
    )

    return (
        spreadsheet_id,
        spreadsheet_name
    )


# ============================================================
# GOOGLE SHEET READ EXISTING DATES
# ============================================================

def get_existing_google_dates(
    sheets,
    spreadsheet_id
):
    """
    Read column A from the station spreadsheet.

    Returns dates such as:

        {
            "2026-09-24",
            "2026-09-25"
        }
    """

    result = (
        sheets.spreadsheets()
        .values()
        .get(
            spreadsheetId=spreadsheet_id,
            range=f"{GOOGLE_SHEET_TAB}!A:A"
        )
        .execute()
    )

    values = result.get(
        "values",
        []
    )

    existing_dates = set()

    for row in values:

        if not row:

            continue

        value = str(
            row[0]
        ).strip()

        if not value:

            continue

        # Skip header.
        if value.lower() == "date":

            continue

        parsed = parse_date_token(
            value
        )

        if parsed:

            existing_dates.add(
                parsed
            )

        else:

            existing_dates.add(
                value
            )

    return existing_dates


# ============================================================
# GOOGLE SHEET APPEND NEW ROWS
# ============================================================

def append_new_google_rows(
    sheets,
    spreadsheet_id,
    entries
):
    """
    Add only dates that are not already present.

    Returns:
        number of newly added rows
    """

    if not entries:

        return 0

    existing_dates = (
        get_existing_google_dates(
            sheets,
            spreadsheet_id
        )
    )

    new_rows = []

    for (
        date_str,
        daily_val
    ) in sorted(entries):

        if date_str in existing_dates:

            continue

        new_rows.append(
            [
                date_str,
                daily_val,
            ]
        )

        # Also update in-memory set in case entries
        # itself contains duplicate dates.
        existing_dates.add(
            date_str
        )

    if not new_rows:

        return 0

    (
        sheets.spreadsheets()
        .values()
        .append(
            spreadsheetId=spreadsheet_id,
            range=f"{GOOGLE_SHEET_TAB}!A:B",
            valueInputOption="USER_ENTERED",
            insertDataOption="INSERT_ROWS",
            body={
                "values": new_rows
            }
        )
        .execute()
    )

    return len(
        new_rows
    )


# ============================================================
# GOOGLE PROCESS ONE STATION
# ============================================================

def save_station_to_google(
    drive,
    sheets,
    folder_id,
    station,
    history
):
    """
    Create/find the station's individual Google Spreadsheet
    and add only new dates.
    """

    (
        spreadsheet_id,
        spreadsheet_name
    ) = get_or_create_station_spreadsheet(
        drive,
        sheets,
        folder_id,
        station
    )

    if not history:

        print(
            "   Google: no scraped rows to upload."
        )

        return 0

    added = append_new_google_rows(
        sheets,
        spreadsheet_id,
        history
    )

    if added:

        print(
            f"   Google: added "
            f"{added} new row(s) to "
            f"{spreadsheet_name}"
        )

    else:

        print(
            "   Google: no new dates. "
            "Everything was already present."
        )

    return added


# ============================================================
# MAIN
# ============================================================

def main():

    # --------------------------------------------------------
    # Create local directories
    # --------------------------------------------------------

    os.makedirs(
        OUTPUT_DIR,
        exist_ok=True
    )

    os.makedirs(
        DEBUG_DIR,
        exist_ok=True
    )

    # --------------------------------------------------------
    # Load stations
    # --------------------------------------------------------

    with open(
        STATIONS_FILE,
        encoding="utf-8",
    ) as f:

        stations = json.load(
            f
        )

    if DEBUG_STATIONS:

        stations = stations[
            :DEBUG_STATIONS
        ]

        print(
            f"DEBUG MODE: running only "
            f"{len(stations)} station(s)"
        )

    else:

        print(
            f"LIVE RUN: processing all "
            f"{len(stations)} station(s) "
            f"in {STATIONS_FILE}."
        )

    # --------------------------------------------------------
    # Google connection
    # --------------------------------------------------------

    print()
    print(
        "======================================"
    )
    print(
        "Connecting to Google Drive / Sheets"
    )
    print(
        "======================================"
    )

    google_creds = (
        load_google_credentials()
    )

    drive = build(
        "drive",
        "v3",
        credentials=google_creds
    )

    sheets = build(
        "sheets",
        "v4",
        credentials=google_creds
    )

    google_folder_id = (
        get_google_drive_folder(
            drive,
            GOOGLE_FOLDER_NAME
        )
    )

    print(
        "Google connection ready."
    )

    # --------------------------------------------------------
    # Run timestamp
    # --------------------------------------------------------

    run_timestamp = (
        datetime.datetime.utcnow()
        .strftime(
            "%Y-%m-%d %H:%M:%S"
        )
    )

    print(
        f"Run timestamp UTC: "
        f"{run_timestamp}"
    )

    # --------------------------------------------------------
    # Playwright
    # --------------------------------------------------------

    with sync_playwright() as p:

        is_github_actions = (
            os.environ.get(
                "GITHUB_ACTIONS",
                ""
            ).lower()
            == "true"
        )

        if is_github_actions:

            print(
                "Browser mode: "
                "HEADLESS (GitHub Actions)"
            )

        else:

            print(
                "Browser mode: "
                "HEADED (local computer)"
            )

        browser = p.chromium.launch(
            headless=is_github_actions,
            slow_mo=300,
        )

        page = browser.new_page()

        # ----------------------------------------------------
        # Process stations
        # ----------------------------------------------------

        total_google_added = 0

        for station in stations:

            url_id = station[
                "url_id"
            ]

            name = station[
                "name"
            ]

            station_index = station.get(
                "station_index"
            )

            print()
            print(
                "======================================"
            )
            print(
                f"Station: {name}"
            )
            print(
                f"Station Index: "
                f"{station_index}"
            )
            print(
                "======================================"
            )

            # ------------------------------------------------
            # Local filename
            # ------------------------------------------------

            filename = safe_filename(
                name,
                station_index,
                url_id,
            )

            filepath = os.path.join(
                OUTPUT_DIR,
                filename,
            )

            # ------------------------------------------------
            # Debug tag
            # ------------------------------------------------

            debug_tag = None

            if DEBUG_STATIONS:

                clean_name_for_tag = re.sub(
                    r"[^\w-]+",
                    "_",
                    name,
                )

                debug_tag = (
                    f"{url_id}_"
                    f"{clean_name_for_tag}"
                )

            # ------------------------------------------------
            # Scrape station
            # ------------------------------------------------

            try:

                history = process_station(
                    page,
                    station,
                    debug_tag=debug_tag,
                )

            except Exception as e:

                print(
                    f"   FAILED: {e}"
                )

                history = []

            # ------------------------------------------------
            # Print scraped data
            # ------------------------------------------------

            print(
                f"   read "
                f"{len(history)} row(s): "
                f"{history}"
            )

            # ------------------------------------------------
            # Local Excel
            # ------------------------------------------------

            if history:

                try:

                    local_added = (
                        append_all_rows(
                            filepath,
                            history
                        )
                    )

                    print(
                        f"   Local Excel: "
                        f"added {local_added} "
                        f"new row(s)"
                    )

                except Exception as e:

                    print(
                        f"   Local Excel FAILED: "
                        f"{e}"
                    )

            else:

                print(
                    "   Local Excel: "
                    "nothing to write"
                )

            # ------------------------------------------------
            # Google Sheets
            # ------------------------------------------------

            try:

                google_added = (
                    save_station_to_google(
                        drive,
                        sheets,
                        google_folder_id,
                        station,
                        history
                    )
                )

                total_google_added += (
                    google_added
                )

            except Exception as e:

                print(
                    f"   Google upload FAILED: "
                    f"{e}"
                )

            time.sleep(
                0.5
            )

        # ----------------------------------------------------
        # Close browser
        # ----------------------------------------------------

        browser.close()

    # --------------------------------------------------------
    # Final summary
    # --------------------------------------------------------

    print()
    print(
        "======================================"
    )
    print(
        "RUN COMPLETED"
    )
    print(
        "======================================"
    )

    print(
        f"Google rows added this run: "
        f"{total_google_added}"
    )

    print(
        f"Output folder: "
        f"{OUTPUT_DIR}"
    )

    print(
        "Done."
    )


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":

    main()
